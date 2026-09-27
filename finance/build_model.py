"""Génère le modèle financier SMARTFIN (5 ans) : finance/SMARTFIN_modele_financier.xlsx

Toutes les projections sont des formules Excel reliées à l'onglet « Hypothèses » :
modifier une cellule bleue recalcule tout le modèle.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).with_name("SMARTFIN_modele_financier.xlsx")
YEARS = ["2027", "2028", "2029", "2030", "2031"]
COLS = ["C", "D", "E", "F", "G"]
HYP = "'Hypothèses'"
PROJ = "Projections"

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14, color="FFFFFF")
SECTION = Font(name=FONT, bold=True, color="2F2C8F")
HEAD_FILL = PatternFill("solid", fgColor="2F2C8F")
SECTION_FILL = PatternFill("solid", fgColor="ECEBFA")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOP = Border(top=Side(style="thin", color="000000"))

MFCFA = '#,##0.0,," M";(#,##0.0,," M");"-"'
FCFA = '#,##0;(#,##0);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
RATIO = '0.0"x";(0.0"x");"-"'

# ───────────────────────── Hypothèses ─────────────────────────
# (clé, libellé, unité, valeur(s), format, commentaire / source, clé ?)
HYPOTHESES = [
    ("§", "Clients"),
    ("users_end", "Inscrits en fin d'année", "clients", [25_000, 90_000, 220_000, 400_000, 650_000], INT,
     "Hypothèse de lancement : Abidjan A1, toute la Côte d'Ivoire A2, Sénégal + Burkina A3, reste UEMOA A4-A5.", True),
    ("active_rate", "Taux d'actifs mensuels (sur inscrits moyens)", "%", [0.35, 0.40, 0.45, 0.48, 0.50], PCT,
     "Part des inscrits ayant au moins une opération dans le mois. Référence fintechs d'épargne : 30-50 %.", True),
    ("plus_rate", "Actifs abonnés SMARTFIN+", "%", [0.03, 0.05, 0.07, 0.08, 0.09], PCT,
     "Conversion freemium typique : 2-10 % des actifs.", True),
    ("§", "Tontines en ligne (produit cœur)"),
    ("tontine_share", "Actifs membres d'au moins une tontine", "%", [0.40, 0.50, 0.55, 0.60, 0.60], PCT,
     "La tontine est le produit d'entrée : la majorité des actifs y cotisent.", True),
    ("tontine_monthly", "Cotisation mensuelle moyenne par membre", "FCFA", [25_000, 27_000, 30_000, 32_000, 35_000], FCFA,
     "Équivalent mensuel (tontines hebdomadaires ramenées au mois). Inclus dans les dépôts.", True),
    ("tontine_commission", "Commission SMARTFIN sur chaque pot versé", "%", [0.01], PCT, "backend/src/pricing.ts → tontine.commissionRate", True),
    ("§", "Microfinance (crédit financé par le SFD partenaire)"),
    ("adv_uptake", "Membres de tontine prenant une avance dans l'année", "%", [0.10, 0.15, 0.20, 0.20, 0.20], PCT,
     "Avance sur le futur pot, remboursée automatiquement au versement.", True),
    ("adv_avg", "Avance moyenne", "FCFA", [40_000, 45_000, 50_000, 55_000, 60_000], FCFA, "Plafond : 50 % du pot.", False),
    ("adv_months", "Durée moyenne d'une avance", "mois", [2], INT, "", False),
    ("adv_rate", "Taux mensuel de l'avance", "%", [0.015], PCT, "18 % / an, sous le taux d'usure SFD (24 %).", False),
    ("mc_uptake", "Actifs ayant un microcrédit dans l'année", "%", [0.01, 0.03, 0.05, 0.07, 0.08], PCT,
     "Réservé aux clients KYC 1+ avec un score ≥ 550.", True),
    ("mc_per_year", "Microcrédits par emprunteur et par an", "prêts", [2], INT, "Prêts de 4 mois en moyenne, renouvelés.", False),
    ("mc_avg", "Microcrédit moyen", "FCFA", [75_000, 90_000, 100_000, 120_000, 140_000], FCFA, "Plafond produit : 500 000 F.", False),
    ("mc_months", "Durée moyenne d'un microcrédit", "mois", [4], INT, "", False),
    ("mc_rate", "Taux mensuel du microcrédit (capital restant dû)", "%", [0.015], PCT, "", False),
    ("smartfin_share", "Part des intérêts revenant à SMARTFIN", "%", [0.40], PCT,
     "Le SFD partenaire finance et porte les prêts ; SMARTFIN note, verse et recouvre. À négocier.", True),
    ("adv_loss", "Pertes sur avances (% des montants versés)", "%", [0.005], PCT, "Faibles : remboursées sur le pot, cautions en plus.", False),
    ("mc_loss", "Pertes sur microcrédits (% des montants versés)", "%", [0.04, 0.035, 0.03, 0.03, 0.03], PCT,
     "Référence microfinance UEMOA : viser un PAR30 < 5 %.", True),
    ("first_loss", "Part des pertes garantie par SMARTFIN (première perte)", "%", [0.50], PCT,
     "Les SFD demandent souvent une garantie de première perte au partenaire technique.", True),
    ("bureau_cost", "Consultation bureau de crédit par prêt", "FCFA", [500], FCFA, "Bureau d'information sur le crédit (BIC) UEMOA.", False),
    ("§", "Flux d'argent"),
    ("dep_month", "Dépôt moyen mensuel par actif", "FCFA", [20_000, 22_000, 25_000, 27_000, 30_000], FCFA,
     "Montant épargné/transité chaque mois via l'app (pas le salaire complet).", True),
    ("wd_share", "Part des dépôts retirée vers Mobile Money", "%", [0.45], PCT,
     "Le reste est épargné (e-caisse, projets, challenges, cagnottes) ou transféré entre utilisateurs.", False),
    ("free_wd_share", "Part des retraits offerts (SMARTFIN+)", "%", [0.10], PCT, "3 retraits offerts par mois aux abonnés.", False),
    ("aum_per_active", "Encours moyen par actif", "FCFA", [25_000, 30_000, 35_000, 40_000, 45_000], FCFA,
     "Solde moyen portefeuille + épargne bloquée.", True),
    ("cagnotte_share", "Part des dépôts passant par une Cagnotte+", "%", [0.15, 0.18, 0.20, 0.20, 0.20], PCT, "Tontines et épargne de groupe.", False),
    ("locked_share", "Part de l'encours en e-caisse", "%", [0.30], PCT, "", False),
    ("break_rate", "Part des e-caisses cassées avant l'échéance (par an)", "%", [0.15], PCT, "", False),
    ("§", "Tarifs (alignés sur backend/src/pricing.ts)"),
    ("wd_fee", "Frais de retrait", "%", [0.015], PCT, "1,5 % (min. 100 F, max. 7 500 F).", True),
    ("cag_fee", "Frais de retrait Cagnotte+", "%", [0.01], PCT, "", False),
    ("break_fee", "Pénalité de rupture e-caisse", "%", [0.02], PCT, "", False),
    ("plus_price", "Abonnement SMARTFIN+ mensuel", "FCFA", [1_000], FCFA, "", True),
    ("float_yield", "Rendement annuel du float (compte de cantonnement / DAT)", "%", [0.03], PCT,
     "À valider avec la banque partenaire : la rémunération des fonds de monnaie électronique est encadrée par la BCEAO.", True),
    ("partner_rev", "Commissions micro-assurance et autres partenaires par actif et par an", "FCFA", [0, 200, 400, 600, 800], FCFA,
     "Apport d'affaires : micro-assurance (santé, décès pour les tontines), épargne rémunérée partenaire.", False),
    ("§", "Coûts directs"),
    ("col_cost", "Coût d'encaissement (agrégateur / opérateurs)", "%", [0.010, 0.009, 0.008, 0.008, 0.007], PCT,
     "Grille indicative agrégateurs UEMOA 1-2 % ; objectif : contrats directs Wave / Orange / MTN < 1 %. HYPOTHÈSE CRITIQUE.", True),
    ("pay_cost", "Coût de décaissement (payout)", "%", [0.006, 0.005, 0.005, 0.004, 0.004], PCT, "À négocier.", True),
    ("sms_month", "SMS par actif et par mois (OTP, alertes)", "SMS", [2], INT, "Réduire via notifications push.", False),
    ("sms_cost", "Coût unitaire SMS", "FCFA", [15], FCFA, "", False),
    ("kyc_rate", "Part des nouveaux inscrits vérifiés (KYC)", "%", [0.40], PCT, "", False),
    ("kyc_cost", "Coût d'une vérification KYC", "FCFA", [300], FCFA, "Vérification pièce + selfie (prestataire).", False),
    ("ref_share", "Part des nouveaux inscrits venus par parrainage", "%", [0.30], PCT, "", False),
    ("ref_bonus", "Bonus parrain", "FCFA", [500], FCFA, "Versé au 1er dépôt ≥ 5 000 F.", False),
    ("§", "Charges d'exploitation"),
    ("headcount", "Effectif (ETP moyen)", "personnes", [8, 14, 22, 32, 42], INT, "Tech, produit, conformité, support, ventes/ambassadeurs.", True),
    ("salary", "Coût moyen annuel chargé par ETP", "FCFA", [7_200_000, 7_560_000, 7_938_000, 8_335_000, 8_752_000], FCFA,
     "600 000 F/mois chargés en A1, +5 %/an.", False),
    ("cac", "Coût d'acquisition par nouvel inscrit (CAC)", "FCFA", [1_500, 1_300, 1_200, 1_100, 1_000], FCFA,
     "Réseaux sociaux, influenceurs, ambassadeurs, activations terrain. Baisse avec le parrainage et la notoriété.", True),
    ("brand", "Marketing de marque (fixe)", "FCFA", [15_000_000, 30_000_000, 50_000_000, 70_000_000, 90_000_000], FCFA, "", False),
    ("cloud_fixed", "Infrastructure cloud (fixe)", "FCFA", [6_000_000, 9_000_000, 15_000_000, 22_000_000, 30_000_000], FCFA, "", False),
    ("cloud_var", "Infrastructure par actif et par an", "FCFA", [120], FCFA, "", False),
    ("partner_fee", "Banque / EME partenaire (frais annuels)", "FCFA", [12_000_000, 15_000_000, 20_000_000, 25_000_000, 30_000_000], FCFA,
     "Tenue du compte de cantonnement et partenariat monnaie électronique.", False),
    ("legal", "Juridique, conformité, audit, sécurité", "FCFA", [10_000_000, 15_000_000, 25_000_000, 30_000_000, 35_000_000], FCFA,
     "Commissaire aux comptes, LBC/FT, tests d'intrusion, assurances.", False),
    ("overhead", "Loyer et frais généraux", "FCFA", [9_000_000, 12_000_000, 18_000_000, 24_000_000, 30_000_000], FCFA, "", False),
    ("tax_rate", "Impôt sur les sociétés", "%", [0.25], PCT, "Taux IS Côte d'Ivoire : 25 % (simplifié : pas de report de déficit).", False),
    ("lifetime", "Durée de vie moyenne d'un client actif", "années", [3], INT, "Pour le calcul de la LTV.", False),
]


def style_header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center" if i > 2 else "left", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE
    for col in range(1, 9):
        ws.cell(row=1, column=col).fill = HEAD_FILL
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONT, italic=True, color="555555")


wb = Workbook()

# ── Onglet Hypothèses ──
hyp = wb.active
hyp.title = "Hypothèses"
title(hyp, "SMARTFIN — Hypothèses du modèle", "Cellules bleues = saisies modifiables · fond jaune = hypothèses clés · montants en FCFA")
hyp["A3"] = "Légende : modifiez uniquement les cellules en bleu ; tous les autres onglets se recalculent. Une valeur seule en colonne C s'applique aux 5 années."
hyp["A3"].font = Font(name=FONT, size=9, color="555555")
style_header(hyp, 5, ["Hypothèse", "Unité", *YEARS, "Commentaire / source"], [52, 11, 14, 14, 14, 14, 14, 80])
ref = {}
r = 6
for item in HYPOTHESES:
    if item[0] == "§":
        hyp.cell(row=r, column=1, value=item[1]).font = SECTION
        for col in range(1, 9):
            hyp.cell(row=r, column=col).fill = SECTION_FILL
        r += 1
        continue
    key, label, unit, values, fmt, note, is_key = item
    hyp.cell(row=r, column=1, value=label).font = BLACK
    hyp.cell(row=r, column=2, value=unit).font = BLACK
    for i, v in enumerate(values):
        c = hyp.cell(row=r, column=3 + i, value=v)
        c.font = BLUE
        c.number_format = fmt
        if is_key:
            c.fill = KEY_FILL
    hyp.cell(row=r, column=8, value=note).font = Font(name=FONT, size=9, color="555555")
    ref[key] = (r, len(values) > 1)
    r += 1
hyp.freeze_panes = "C6"


def h(key, col):
    """Référence à une hypothèse pour la colonne d'année `col`."""
    row, per_year = ref[key]
    return f"{HYP}!{col}{row}" if per_year else f"{HYP}!$C${row}"


# ── Onglet Investissement ──
inv = wb.create_sheet("Investissement")
title(inv, "SMARTFIN — Budget de lancement", "Investissements et dépenses avant/au démarrage (FCFA)")
style_header(inv, 4, ["Poste", "Montant (FCFA)", "Commentaire"], [52, 18, 80])
INVEST = [
    ("Développement app mobile + API + back-office (finitions, audits)", 25_000_000, "Base technique livrée dans ce dépôt ; budget de durcissement et d'intégrations opérateurs."),
    ("Design UX/UI, identité, contenus tutoriels vidéo", 6_000_000, ""),
    ("Test d'intrusion et audit de sécurité avant lancement", 5_000_000, "Obligatoire avant d'ouvrir les flux d'argent."),
    ("Constitution société, contrats partenaires, conseil réglementaire", 8_000_000, "Convention avec banque / EME agréé BCEAO, CGU, politique LBC/FT."),
    ("Équipement (ordinateurs, téléphones de test)", 6_000_000, "Amorti sur 3 ans."),
    ("Campagne de lancement (événement, influenceurs, affichage)", 12_000_000, ""),
]
for i, (label, amount, note) in enumerate(INVEST, start=5):
    inv.cell(row=i, column=1, value=label).font = BLACK
    c = inv.cell(row=i, column=2, value=amount)
    c.font = BLUE
    c.number_format = FCFA
    inv.cell(row=i, column=3, value=note).font = Font(name=FONT, size=9, color="555555")
last = 4 + len(INVEST)
inv.cell(row=last + 1, column=1, value="Total investissements de lancement").font = BOLD
tot = inv.cell(row=last + 1, column=2, value=f"=SUM(B5:B{last})")
tot.font = BOLD
tot.number_format = FCFA
tot.border = TOP
CAPEX_REF = f"Investissement!$B${last + 1}"
inv.cell(row=last + 2, column=1, value="Durée d'amortissement (années)").font = BLACK
dur = inv.cell(row=last + 2, column=2, value=3)
dur.font = BLUE
AMORT_YEARS = f"Investissement!$B${last + 2}"

# ── Onglet Projections ──
p = wb.create_sheet(PROJ)
title(p, "SMARTFIN — Projections sur 5 ans (tontine en ligne + microfinance)", "Montants en millions de FCFA sauf mention contraire · tout est calculé à partir des hypothèses")
style_header(p, 4, ["Indicateur", "Unité", *YEARS, "Méthode de calcul"], [46, 12, 14, 14, 14, 14, 14, 60])
p.freeze_panes = "C5"
rows = {}
cur = 5


def section(label):
    global cur
    p.cell(row=cur, column=1, value=label).font = SECTION
    for col in range(1, 9):
        p.cell(row=cur, column=col).fill = SECTION_FILL
    cur += 1


def line(key, label, unit, formula, fmt=MFCFA, note="", bold=False, total=False):
    """`formula(col, prev_col)` renvoie la formule de l'année ; prev_col vaut None la première année."""
    global cur
    rows[key] = cur  # enregistré d'abord : une ligne cumulative se référence elle-même (année précédente)
    p.cell(row=cur, column=1, value=label).font = BOLD if bold else BLACK
    p.cell(row=cur, column=2, value=unit).font = BLACK
    for i, col in enumerate(COLS):
        c = p[f"{col}{cur}"]
        c.value = formula(col, COLS[i - 1] if i else None)
        c.number_format = fmt
        c.font = BOLD if bold else (GREEN if c.value.startswith(f"={HYP}") or c.value.startswith("=Investissement") else BLACK)
        if total:
            c.border = TOP
    p.cell(row=cur, column=8, value=note).font = Font(name=FONT, size=9, color="555555")
    cur += 1


R = lambda key, col: f"{col}{rows[key]}"  # noqa: E731

section("Clients")
line("users_end", "Inscrits en fin d'année", "clients", lambda c, pc: f"={h('users_end', c)}", INT)
line("users_avg", "Inscrits moyens", "clients",
     lambda c, pc: f"=({pc}{rows['users_end']}+{c}{rows['users_end']})/2" if pc else f"={c}{rows['users_end']}/2", INT, "(début + fin) / 2")
line("new_users", "Nouveaux inscrits", "clients",
     lambda c, pc: f"={c}{rows['users_end']}-{pc}{rows['users_end']}" if pc else f"={c}{rows['users_end']}", INT)
line("actives", "Clients actifs (moyenne mensuelle)", "clients", lambda c, pc: f"={R('users_avg', c)}*{h('active_rate', c)}", INT)
line("plus_subs", "Abonnés SMARTFIN+", "clients", lambda c, pc: f"={R('actives', c)}*{h('plus_rate', c)}", INT)

section("Volumes")
line("deposits", "Dépôts annuels", "M FCFA", lambda c, pc: f"={R('actives', c)}*{h('dep_month', c)}*12", note="actifs × dépôt mensuel × 12")
line("withdrawals", "Retraits vers Mobile Money", "M FCFA", lambda c, pc: f"={R('deposits', c)}*{h('wd_share', c)}")
line("paid_wd", "dont retraits payants", "M FCFA", lambda c, pc: f"={R('withdrawals', c)}*(1-{h('free_wd_share', c)})")
line("aum", "Encours moyen clients (float)", "M FCFA", lambda c, pc: f"={R('actives', c)}*{h('aum_per_active', c)}")
line("cag_wd", "Retraits de Cagnottes+", "M FCFA", lambda c, pc: f"={R('deposits', c)}*{h('cagnotte_share', c)}")
line("breaks", "E-caisses cassées avant l'échéance", "M FCFA", lambda c, pc: f"={R('aum', c)}*{h('locked_share', c)}*{h('break_rate', c)}")

line("tontiniers", "Membres de tontine actifs", "clients", lambda c, pc: f"={R('actives', c)}*{h('tontine_share', c)}", INT)
line("tontine_vol", "Pots de tontine versés", "M FCFA", lambda c, pc: f"={R('tontiniers', c)}*{h('tontine_monthly', c)}*12",
     note="membres × cotisation mensuelle × 12 (chaque cotisation finit dans un pot)")
line("adv_vol", "Avances sur tontine versées", "M FCFA", lambda c, pc: f"={R('tontiniers', c)}*{h('adv_uptake', c)}*{h('adv_avg', c)}")
line("mc_count", "Microcrédits accordés", "prêts", lambda c, pc: f"={R('actives', c)}*{h('mc_uptake', c)}*{h('mc_per_year', c)}", INT)
line("mc_vol", "Microcrédits versés", "M FCFA", lambda c, pc: f"={R('mc_count', c)}*{h('mc_avg', c)}")
line("credit_total", "Total crédits versés", "M FCFA", lambda c, pc: f"={R('adv_vol', c)}+{R('mc_vol', c)}")
line("credit_int", "Intérêts totaux (SMARTFIN + SFD)", "M FCFA",
     lambda c, pc: f"={R('adv_vol', c)}*{h('adv_rate', c)}*{h('adv_months', c)}+{R('mc_vol', c)}*{h('mc_rate', c)}*({h('mc_months', c)}+1)/2",
     note="avances : taux × durée ; microcrédit amorti : taux × (n+1)/2")

section("Revenus")
line("rev_tontine", "Commissions sur pots de tontine", "M FCFA", lambda c, pc: f"={R('tontine_vol', c)}*{h('tontine_commission', c)}")
line("rev_credit", "Part SMARTFIN des intérêts de crédit", "M FCFA", lambda c, pc: f"={R('credit_int', c)}*{h('smartfin_share', c)}")
line("rev_wd", "Frais de retrait", "M FCFA", lambda c, pc: f"={R('paid_wd', c)}*{h('wd_fee', c)}")
line("rev_cag", "Frais Cagnotte+", "M FCFA", lambda c, pc: f"={R('cag_wd', c)}*{h('cag_fee', c)}")
line("rev_break", "Pénalités e-caisse", "M FCFA", lambda c, pc: f"={R('breaks', c)}*{h('break_fee', c)}")
line("rev_plus", "Abonnements SMARTFIN+", "M FCFA", lambda c, pc: f"={R('plus_subs', c)}*{h('plus_price', c)}*12")
line("rev_float", "Revenus du float", "M FCFA", lambda c, pc: f"={R('aum', c)}*{h('float_yield', c)}")
line("rev_partner", "Commissions partenaires", "M FCFA", lambda c, pc: f"={R('actives', c)}*{h('partner_rev', c)}")
line("revenue", "Chiffre d'affaires", "M FCFA", lambda c, pc: f"=SUM({c}{rows['rev_tontine']}:{c}{rows['rev_partner']})", bold=True, total=True)

section("Coûts directs")
line("c_col", "Frais d'encaissement opérateurs", "M FCFA", lambda c, pc: f"={R('deposits', c)}*{h('col_cost', c)}")
line("c_pay", "Frais de décaissement opérateurs", "M FCFA", lambda c, pc: f"={R('withdrawals', c)}*{h('pay_cost', c)}")
line("c_sms", "SMS (OTP, alertes)", "M FCFA", lambda c, pc: f"={R('actives', c)}*{h('sms_month', c)}*12*{h('sms_cost', c)}")
line("c_kyc", "Vérifications KYC", "M FCFA", lambda c, pc: f"={R('new_users', c)}*{h('kyc_rate', c)}*{h('kyc_cost', c)}")
line("c_ref", "Bonus de parrainage", "M FCFA", lambda c, pc: f"={R('new_users', c)}*{h('ref_share', c)}*{h('ref_bonus', c)}")
line("c_risk", "Coût du risque (garantie de première perte)", "M FCFA",
     lambda c, pc: f"=({R('adv_vol', c)}*{h('adv_loss', c)}+{R('mc_vol', c)}*{h('mc_loss', c)})*{h('first_loss', c)}")
line("c_bureau", "Bureau de crédit", "M FCFA",
     lambda c, pc: f"=({R('mc_count', c)}+{R('tontiniers', c)}*{h('adv_uptake', c)})*{h('bureau_cost', c)}")
line("direct", "Total coûts directs", "M FCFA", lambda c, pc: f"=SUM({c}{rows['c_col']}:{c}{rows['c_bureau']})", bold=True, total=True)
line("gross", "Marge brute", "M FCFA", lambda c, pc: f"={R('revenue', c)}-{R('direct', c)}", bold=True)
line("gross_pct", "Taux de marge brute", "%", lambda c, pc: f"=IF({R('revenue', c)}=0,0,{R('gross', c)}/{R('revenue', c)})", PCT)

section("Charges d'exploitation")
line("o_staff", "Salaires et charges", "M FCFA", lambda c, pc: f"={h('headcount', c)}*{h('salary', c)}")
line("o_acq", "Marketing d'acquisition", "M FCFA", lambda c, pc: f"={R('new_users', c)}*{h('cac', c)}", note="nouveaux inscrits × CAC")
line("o_brand", "Marketing de marque", "M FCFA", lambda c, pc: f"={h('brand', c)}")
line("o_cloud", "Infrastructure cloud", "M FCFA", lambda c, pc: f"={h('cloud_fixed', c)}+{R('actives', c)}*{h('cloud_var', c)}")
line("o_partner", "Banque / EME partenaire", "M FCFA", lambda c, pc: f"={h('partner_fee', c)}")
line("o_legal", "Juridique, conformité, sécurité", "M FCFA", lambda c, pc: f"={h('legal', c)}")
line("o_overhead", "Loyer et frais généraux", "M FCFA", lambda c, pc: f"={h('overhead', c)}")
line("opex", "Total charges d'exploitation", "M FCFA", lambda c, pc: f"=SUM({c}{rows['o_staff']}:{c}{rows['o_overhead']})", bold=True, total=True)

section("Résultat")
line("ebitda", "EBITDA", "M FCFA", lambda c, pc: f"={R('gross', c)}-{R('opex', c)}", bold=True)
line("ebitda_pct", "Marge d'EBITDA", "%", lambda c, pc: f"=IF({R('revenue', c)}=0,0,{R('ebitda', c)}/{R('revenue', c)})", PCT)
line("amort", "Amortissements", "M FCFA",
     lambda c, pc: f"=IF({COLS.index(c) + 1}<={AMORT_YEARS},{CAPEX_REF}/{AMORT_YEARS},0)", note="Investissements de lancement amortis linéairement")
line("ebit", "Résultat d'exploitation", "M FCFA", lambda c, pc: f"={R('ebitda', c)}-{R('amort', c)}")
line("tax", "Impôt sur les sociétés", "M FCFA", lambda c, pc: f"=MAX(0,{R('ebit', c)})*{h('tax_rate', c)}")
line("net", "Résultat net", "M FCFA", lambda c, pc: f"={R('ebit', c)}-{R('tax', c)}", bold=True, total=True)

section("Trésorerie")
line("capex", "Investissements de lancement", "M FCFA", lambda c, pc: f"=-{CAPEX_REF}" if c == "C" else "=0")
line("cash_flow", "Flux de trésorerie", "M FCFA", lambda c, pc: f"={R('net', c)}+{R('amort', c)}+{R('capex', c)}",
     note="Résultat net + amortissements − investissements (BFR négligé : les fonds clients sont cantonnés)")
line("cash_cum", "Trésorerie cumulée (hors levée)", "M FCFA",
     lambda c, pc: f"={pc}{rows['cash_cum']}+{R('cash_flow', c)}" if pc else f"={R('cash_flow', c)}", bold=True)
line("ebitda_pos", "EBITDA positif ? (1 = oui)", "indicateur", lambda c, pc: f"=IF({R('ebitda', c)}>0,1,0)", INT)

section("Économie unitaire (par client actif et par an, en FCFA)")
line("u_rev", "Revenu par actif", "FCFA", lambda c, pc: f"=IF({R('actives', c)}=0,0,{R('revenue', c)}/{R('actives', c)})", FCFA)
line("u_direct", "Coûts directs par actif", "FCFA", lambda c, pc: f"=IF({R('actives', c)}=0,0,{R('direct', c)}/{R('actives', c)})", FCFA)
line("u_margin", "Marge brute par actif", "FCFA", lambda c, pc: f"={R('u_rev', c)}-{R('u_direct', c)}", FCFA)
line("u_cac_active", "Coût d'acquisition par client actif", "FCFA",
     lambda c, pc: f"=IF({h('active_rate', c)}=0,0,{h('cac', c)}/{h('active_rate', c)})", FCFA, "CAC par inscrit ÷ taux d'actifs")
line("u_ltv", "Valeur vie client (LTV)", "FCFA", lambda c, pc: f"={R('u_margin', c)}*{h('lifetime', c)}", FCFA, "marge brute par actif × durée de vie")
line("u_ltv_cac", "Ratio LTV / CAC", "x", lambda c, pc: f"=IF({R('u_cac_active', c)}=0,0,{R('u_ltv', c)}/{R('u_cac_active', c)})", RATIO, "Objectif > 3x")

# ── Onglet Synthèse ──
s = wb.create_sheet("Synthèse", 0)
title(s, "SMARTFIN — Synthèse financière", "Vue d'ensemble pour investisseurs et partenaires (M FCFA sauf mention contraire)")
style_header(s, 4, ["Indicateur clé", "Unité", *YEARS], [40, 12, 14, 14, 14, 14, 14])
SYN = [
    ("Inscrits en fin d'année", "clients", "users_end", INT),
    ("Clients actifs", "clients", "actives", INT),
    ("Dépôts annuels", "M FCFA", "deposits", MFCFA),
    ("Encours moyen clients", "M FCFA", "aum", MFCFA),
    ("Pots de tontine versés", "M FCFA", "tontine_vol", MFCFA),
    ("Crédits versés (avances + microcrédits)", "M FCFA", "credit_total", MFCFA),
    ("Chiffre d'affaires", "M FCFA", "revenue", MFCFA),
    ("Marge brute", "M FCFA", "gross", MFCFA),
    ("Taux de marge brute", "%", "gross_pct", PCT),
    ("EBITDA", "M FCFA", "ebitda", MFCFA),
    ("Résultat net", "M FCFA", "net", MFCFA),
    ("Trésorerie cumulée (hors levée)", "M FCFA", "cash_cum", MFCFA),
    ("Revenu par actif", "FCFA", "u_rev", FCFA),
    ("Ratio LTV / CAC", "x", "u_ltv_cac", RATIO),
]
for i, (label, unit, key, fmt) in enumerate(SYN, start=5):
    s.cell(row=i, column=1, value=label).font = BLACK
    s.cell(row=i, column=2, value=unit).font = BLACK
    for col in COLS:
        c = s[f"{col}{i}"]
        c.value = f"={PROJ}!{col}{rows[key]}"
        c.font = GREEN
        c.number_format = fmt
r0 = 5 + len(SYN) + 1
s.cell(row=r0, column=1, value="Chiffres clés").font = SECTION
for col in range(1, 8):
    s.cell(row=r0, column=col).fill = SECTION_FILL
cc = lambda key: f"{PROJ}!$C${rows[key]}:$G${rows[key]}"  # noqa: E731
KEYS = [
    ("Besoin de financement (creux de trésorerie)", f"=-MIN(0,MIN({cc('cash_cum')}))", MFCFA,
     "Montant minimal à lever avant marge de sécurité ; prévoir +30 %."),
    ("Levée recommandée (besoin + 30 % de sécurité)", f"=B{r0 + 1}*1.3", MFCFA, ""),
    ("Première année d'EBITDA positif", f'=IFERROR(INDEX({PROJ}!$C$4:$G$4,MATCH(1,{cc("ebitda_pos")},0)),"Après 2031")', "@", ""),
    ("Chiffre d'affaires cumulé sur 5 ans", f"=SUM({cc('revenue')})", MFCFA, ""),
    ("Investissements de lancement", f"={CAPEX_REF}", MFCFA, ""),
]
for i, (label, formula, fmt, note) in enumerate(KEYS, start=r0 + 1):
    s.cell(row=i, column=1, value=label).font = BOLD
    c = s.cell(row=i, column=2, value=formula)
    c.number_format = fmt
    c.font = BOLD
    s.cell(row=i, column=3, value=note).font = Font(name=FONT, size=9, color="555555")
s.column_dimensions["B"].width = 16
s[f"A{r0 + len(KEYS) + 2}"] = ("Lecture : SMARTFIN est une tontine en ligne adossée à la microfinance. L'équilibre dépend surtout (1) du coût d'encaissement "
                               "négocié avec les opérateurs, (2) de la part des actifs membres de tontines et (3) de la part des intérêts "
                               "négociée avec le SFD partenaire, sous contrainte du coût du risque. "
                               "Voir l'onglet Sensibilité et les cellules jaunes de l'onglet Hypothèses.")
s[f"A{r0 + len(KEYS) + 2}"].font = Font(name=FONT, size=9, italic=True)
s[f"A{r0 + len(KEYS) + 2}"].alignment = Alignment(wrap_text=True, vertical="top")
s.merge_cells(start_row=r0 + len(KEYS) + 2, start_column=1, end_row=r0 + len(KEYS) + 2, end_column=7)
s.row_dimensions[r0 + len(KEYS) + 2].height = 42

# ── Onglet Sensibilité : EBITDA 2029 selon coût d'encaissement × frais de retrait ──
sens = wb.create_sheet("Sensibilité")
title(sens, "SMARTFIN — Sensibilité de l'EBITDA 2029 (M FCFA)", "Coût d'encaissement (lignes) × frais de retrait (colonnes), autres hypothèses inchangées")
sens.column_dimensions["A"].width = 34
for col in "BCDEF":
    sens.column_dimensions[col].width = 14
sens["A4"] = "Coût d'encaissement ↓  /  Frais de retrait →"
sens["A4"].font = BOLD
fees = [0.01, 0.0125, 0.015, 0.0175, 0.02]
costs = [0.005, 0.0065, 0.008, 0.01, 0.015]
for j, f in enumerate(fees):
    c = sens.cell(row=4, column=2 + j, value=f)
    c.font = Font(name=FONT, bold=True, color="0000FF")
    c.number_format = PCT
# EBITDA 2029 recalculé analytiquement à partir des lignes de Projections (colonne E = 2029) :
# EBITDA = EBITDA_base − Δfrais_retrait − Δcoût_encaissement
E = "E"
base_ebitda = f"{PROJ}!${E}${rows['ebitda']}"
paid_wd = f"{PROJ}!${E}${rows['paid_wd']}"
deposits = f"{PROJ}!${E}${rows['deposits']}"
wd_fee_ref = h("wd_fee", E)
col_cost_ref = f"{HYP}!$E${ref['col_cost'][0]}"
for i, cst in enumerate(costs):
    rr = 5 + i
    c = sens.cell(row=rr, column=1, value=cst)
    c.font = Font(name=FONT, bold=True, color="0000FF")
    c.number_format = PCT
    for j in range(len(fees)):
        col = get_column_letter(2 + j)
        cell = sens[f"{col}{rr}"]
        cell.value = f"={base_ebitda}+{paid_wd}*({col}$4-{wd_fee_ref})-{deposits}*($A{rr}-{col_cost_ref})"
        cell.number_format = MFCFA
        cell.font = BLACK
sens["A11"] = "Les valeurs de l'intersection surlignée correspondent au scénario central de l'onglet Projections."
sens["A11"].font = Font(name=FONT, size=9, italic=True)
sens["D7"].fill = KEY_FILL

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[4 if ws.title != "Hypothèses" else 5].height = 30

hyp[f"C{ref['col_cost'][0]}"].comment = Comment("Hypothèse la plus sensible du modèle : voir l'onglet Sensibilité.", "SMARTFIN")
wb.save(OUT)
print(f"Modèle écrit : {OUT}")
